"""Import de builds depuis un tableur (Excel .xlsx ou CSV).

Une ligne du fichier = un build. Les colonnes portent les noms d'objets et de
sorts tels qu'ils apparaissent dans le catalogue Albion (« Masse », « Casque de
soldat »...). Les cellules de sorts laissees vides sont completees comme dans le
formulaire : sort impose par l'objet, sort natif d'une armure, sinon premier
choix de sa categorie.

Les lignes obtenues repassent par backend/validation.py : un fichier ne peut pas
introduire un build que le site refuserait.

openpyxl est facultatif : sans lui, seul le CSV est accepte et le modele est
fourni au format CSV.
"""
from __future__ import annotations

import csv
import io
import unicodedata

from .catalogue import REGLES, IndexCatalogue
from .models import SlotEquipement
from .validation import ErreurLigne, LIBELLES_CHAMPS, LIBELLES_SLOTS, defauts_pour_objet, valider_lignes

try:  # openpyxl est facultatif : sans lui, l'import se limite au CSV.
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    XLSX_DISPONIBLE = True
except ImportError:  # pragma: no cover - depend de l'installation
    XLSX_DISPONIBLE = False

COLONNE_JOUEUR = "joueur"
ALIAS_JOUEUR = ("joueur", "role", "role_ou_joueur", "nom", "build")
PREMIERE_LIGNE = 2      # ligne 1 = en-tetes
MAX_LIGNES = 200        # garde-fou : un fichier ne cree pas 10 000 builds
FEUILLE_BUILDS = "Builds"
FEUILLE_CATALOGUE = "Catalogue"


def _normaliser(texte: str) -> str:
    """Compare les noms sans se soucier de la casse, des accents ni des espaces."""
    sans_accents = "".join(
        c for c in unicodedata.normalize("NFD", str(texte or ""))
        if unicodedata.category(c) != "Mn"
    )
    return " ".join(sans_accents.lower().replace("_", " ").split())


# --------------------------------------------------------------------------
# Colonnes attendues
# --------------------------------------------------------------------------


def colonnes() -> list[str]:
    """En-tetes du fichier, dans l'ordre : un objet puis ses sorts, slot par slot.

    Les sorts imposes (sort 3 d'une arme, passif d'une cape, sort d'une monture)
    n'ont pas de colonne : ils viennent de l'objet.
    """
    noms = [COLONNE_JOUEUR]
    for slot, regle in REGLES.items():
        noms.append(slot.value)
        noms.extend(f"{slot.value}_{champ.nom}" for champ in regle.champs)
    return noms


def _cible_de_colonne() -> dict[str, tuple[SlotEquipement | None, str | None]]:
    """Nom normalise d'une colonne -> (slot, champ de sort) ; (None, None) = joueur."""
    cibles: dict[str, tuple[SlotEquipement | None, str | None]] = {
        _normaliser(alias): (None, None) for alias in ALIAS_JOUEUR
    }
    for slot, regle in REGLES.items():
        cibles[_normaliser(slot.value)] = (slot, None)
        cibles[_normaliser(LIBELLES_SLOTS[slot])] = (slot, None)
        for champ in regle.champs:
            cibles[_normaliser(f"{slot.value}_{champ.nom}")] = (slot, champ.nom)
    return cibles


# --------------------------------------------------------------------------
# Lecture du fichier
# --------------------------------------------------------------------------


def lire_tableur(contenu: bytes, nom_fichier: str) -> list[list[str]]:
    """Renvoie les cellules du fichier, en-tetes compris, sous forme de texte."""
    extension = nom_fichier.lower().rsplit(".", 1)[-1] if "." in nom_fichier else ""
    if extension in ("csv", "txt"):
        return _lire_csv(contenu)
    if extension in ("xlsx", "xlsm"):
        return _lire_xlsx(contenu)
    raise ErreurLigne([{
        "champ": "fichier",
        "message": "Format non reconnu : attendu .xlsx ou .csv."
                   + ("" if XLSX_DISPONIBLE else " (openpyxl n'est pas installe : "
                      "seul le CSV est accepte sur ce serveur)"),
    }])


def _lire_csv(contenu: bytes) -> list[list[str]]:
    try:
        texte = contenu.decode("utf-8-sig")
    except UnicodeDecodeError:
        texte = contenu.decode("latin-1")
    try:
        dialecte = csv.Sniffer().sniff(texte[:4096], delimiters=";,\t")
    except csv.Error:
        dialecte = csv.excel
        dialecte.delimiter = ";"
    return [[str(cellule) for cellule in ligne] for ligne in csv.reader(io.StringIO(texte), dialecte)]


def _lire_xlsx(contenu: bytes) -> list[list[str]]:
    if not XLSX_DISPONIBLE:
        raise ErreurLigne([{
            "champ": "fichier",
            "message": "Lecture des .xlsx indisponible : installez openpyxl "
                       "(pip install -r requirements.txt), ou enregistrez le fichier en CSV.",
        }])
    try:
        classeur = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    except Exception as erreur:  # openpyxl leve des exceptions variees
        raise ErreurLigne([{
            "champ": "fichier",
            "message": f"Fichier Excel illisible : {erreur}",
        }]) from erreur
    feuille = classeur[FEUILLE_BUILDS] if FEUILLE_BUILDS in classeur.sheetnames else classeur.active
    lignes = [
        ["" if cellule is None else str(cellule).strip() for cellule in rangee]
        for rangee in feuille.iter_rows(values_only=True)
    ]
    classeur.close()
    return lignes


# --------------------------------------------------------------------------
# Conversion en lignes de compo
# --------------------------------------------------------------------------


def _index_objets(index: IndexCatalogue) -> dict[tuple[SlotEquipement, str], int]:
    return {(objet.slot, _normaliser(objet.nom)): objet.id for objet in index.objets.values()}


def importer_lignes(index: IndexCatalogue, contenu: bytes, nom_fichier: str) -> tuple[list[dict], list[str]]:
    """Lit le fichier et renvoie (lignes de compo validees, avertissements).

    Leve ErreurLigne — donc un 422 au format habituel du frontend — des qu'une
    cellule ne correspond a rien de connu ou qu'un slot obligatoire manque.
    """
    if index.est_vide():
        raise ErreurLigne([{
            "champ": "catalogue",
            "message": "Le catalogue Albion n'est pas charge : "
                       "lancez `python manage.py importer-catalogue`.",
        }])

    tableau = lire_tableur(contenu, nom_fichier)
    tableau = [rangee for rangee in tableau if any(str(c).strip() for c in rangee)]
    if not tableau:
        raise ErreurLigne([{"champ": "fichier", "message": "Le fichier est vide."}])

    cibles = _cible_de_colonne()
    entetes = [_normaliser(cellule) for cellule in tableau[0]]
    inconnues = [
        tableau[0][position] for position, entete in enumerate(entetes)
        if entete and entete not in cibles
    ]
    if not any(entete in cibles for entete in entetes):
        raise ErreurLigne([{
            "champ": "fichier",
            "message": "Aucune colonne reconnue en premiere ligne. Attendu : "
                       + ", ".join(colonnes()) + ".",
        }])

    avertissements: list[str] = []
    if inconnues:
        avertissements.append("Colonnes ignorées : " + ", ".join(inconnues) + ".")

    corps = tableau[1:]
    if len(corps) > MAX_LIGNES:
        avertissements.append(f"Seules les {MAX_LIGNES} premières lignes ont été lues.")
        corps = corps[:MAX_LIGNES]

    objets_par_nom = _index_objets(index)
    erreurs: list[dict[str, str]] = []
    lignes: list[dict] = []

    for position, rangee in enumerate(corps):
        numero = position + PREMIERE_LIGNE
        cellules: dict[tuple[SlotEquipement | None, str | None], str] = {}
        for colonne, entete in enumerate(entetes):
            cible = cibles.get(entete)
            if cible is None:
                continue
            cellules[cible] = str(rangee[colonne]).strip() if colonne < len(rangee) else ""

        def signaler(colonne: str, message: str) -> None:
            erreurs.append({"champ": f"Fichier ligne {numero} · {colonne}", "message": message})

        donnees: dict = {"role_ou_joueur": cellules.get((None, None), "")[:120]}

        for slot, regle in REGLES.items():
            nom_objet = cellules.get((slot, None), "")
            libelle = LIBELLES_SLOTS[slot]
            if not nom_objet:
                if regle.objet_obligatoire:
                    signaler(libelle, "objet obligatoire.")
                continue
            objet_id = objets_par_nom.get((slot, _normaliser(nom_objet)))
            if objet_id is None:
                signaler(libelle, f"« {nom_objet} » ne figure pas dans le catalogue "
                                  f"pour ce slot.")
                continue

            donnees[regle.champ_objet()] = objet_id
            donnees.update(defauts_pour_objet(index, objet_id))
            objet = index.objets[objet_id]

            for champ in regle.champs:
                nom_champ = regle.champ_sort(champ.nom)
                autorises = index.pool(objet.categorie_id, champ.emplacement)
                intitule = f"{libelle} · {LIBELLES_CHAMPS.get(champ.nom, champ.nom)}"
                voulu = cellules.get((slot, champ.nom), "")
                if voulu:
                    trouve = next(
                        (identifiant for identifiant in autorises
                         if _normaliser(index.sorts[identifiant].nom) == _normaliser(voulu)),
                        None,
                    )
                    if trouve is None:
                        signaler(intitule, f"« {voulu} » n'est pas disponible pour "
                                           f"« {objet.nom} ».")
                        continue
                    donnees[nom_champ] = trouve
                elif donnees.get(nom_champ) is None and champ.obligatoire and autorises:
                    # Meme repli que le formulaire : le premier choix de la categorie.
                    donnees[nom_champ] = autorises[0]

        lignes.append(donnees)

    if erreurs:
        raise ErreurLigne(erreurs)
    if not lignes:
        raise ErreurLigne([{"champ": "fichier", "message": "Aucun build dans le fichier."}])

    # Derniere passe : les memes regles que pour une saisie a la main.
    try:
        validees = valider_lignes(index, lignes)
    except ErreurLigne as erreur:
        raise ErreurLigne([_reetiqueter(detail) for detail in erreur.erreurs]) from None
    return validees, avertissements


def _reetiqueter(erreur: dict[str, str]) -> dict[str, str]:
    """« lignes.0.offhand_id » -> « Fichier ligne 2 · offhand »."""
    morceaux = erreur.get("champ", "").split(".")
    if morceaux[0] != "lignes" or len(morceaux) < 3:
        return erreur
    champ = ".".join(morceaux[2:]).removesuffix("_id")
    numero = int(morceaux[1]) + PREMIERE_LIGNE
    return {"champ": f"Fichier ligne {numero} · {champ}", "message": erreur["message"]}


# --------------------------------------------------------------------------
# Modele a remplir
# --------------------------------------------------------------------------


def _exemple(index: IndexCatalogue) -> dict[str, str]:
    """Une ligne d'exemple, avec de vrais noms du catalogue et un build coherent."""
    exemple = {COLONNE_JOUEUR: "Tank principal"}
    arme_a_une_main = False
    for slot in REGLES:
        objets = sorted(
            (o for o in index.objets.values() if o.slot == slot), key=lambda o: o.nom
        )
        if not objets:
            continue
        if slot == SlotEquipement.arme:
            # Une arme a une main : le modele peut montrer une off-hand sans se contredire.
            objet = next((o for o in objets if not o.deux_mains), objets[0])
            arme_a_une_main = not objet.deux_mains
        elif slot == SlotEquipement.offhand and not arme_a_une_main:
            continue
        else:
            objet = objets[0]
        exemple[slot.value] = objet.nom
    return exemple


def _noms_par_slot(index: IndexCatalogue) -> dict[SlotEquipement, list[str]]:
    noms: dict[SlotEquipement, list[str]] = {slot: [] for slot in REGLES}
    for objet in index.objets.values():
        if objet.slot in noms:
            noms[objet.slot].append(objet.nom)
    return {slot: sorted(set(liste)) for slot, liste in noms.items()}


def modele_csv(index: IndexCatalogue) -> bytes:
    tampon = io.StringIO()
    ecrivain = csv.writer(tampon, delimiter=";")
    entetes = colonnes()
    ecrivain.writerow(entetes)
    exemple = _exemple(index)
    ecrivain.writerow([exemple.get(colonne, "") for colonne in entetes])
    # BOM : Excel ouvre alors le CSV en UTF-8 sans abimer les accents.
    return tampon.getvalue().encode("utf-8-sig")


def modele_xlsx(index: IndexCatalogue) -> bytes:
    """Classeur pret a remplir : une feuille de saisie + les noms valides en listes."""
    classeur = Workbook()
    builds = classeur.active
    builds.title = FEUILLE_BUILDS
    reference = classeur.create_sheet(FEUILLE_CATALOGUE)

    entetes = colonnes()
    builds.append(entetes)
    for cellule in builds[1]:
        cellule.font = Font(bold=True)
    exemple = _exemple(index)
    builds.append([exemple.get(colonne, "") for colonne in entetes])
    builds.freeze_panes = "A2"
    for position, entete in enumerate(entetes, start=1):
        builds.column_dimensions[get_column_letter(position)].width = max(14, len(entete) + 2)

    # Feuille de reference : un nom d'objet par ligne, une colonne par slot.
    noms = _noms_par_slot(index)
    plages: dict[SlotEquipement, str] = {}
    for position, (slot, liste) in enumerate(noms.items(), start=1):
        lettre = get_column_letter(position)
        reference.cell(row=1, column=position, value=LIBELLES_SLOTS[slot]).font = Font(bold=True)
        for rangee, nom in enumerate(liste, start=2):
            reference.cell(row=rangee, column=position, value=nom)
        reference.column_dimensions[lettre].width = 28
        if liste:
            plages[slot] = f"'{FEUILLE_CATALOGUE}'!${lettre}$2:${lettre}${len(liste) + 1}"

    # Menus deroulants sur les colonnes d'objets : plus de faute de frappe.
    for slot, plage in plages.items():
        colonne = get_column_letter(entetes.index(slot.value) + 1)
        controle = DataValidation(type="list", formula1=plage, allow_blank=True)
        controle.error = "Choisissez un objet de la feuille « Catalogue »."
        builds.add_data_validation(controle)
        controle.add(f"{colonne}2:{colonne}{MAX_LIGNES + 1}")

    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()
